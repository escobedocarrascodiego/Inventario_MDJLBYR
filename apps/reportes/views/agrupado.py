"""Reportes de bienes agrupados por criterio (PDF y Excel).

En vez de escribir una vista por cada reporte pedido (por orden de compra, por
cuenta contable, por denominación, por marca...) y duplicar el mismo armado de
tabla ocho veces, hay UNA vista parametrizada:

    /funciones/reportes/agrupado/?agrupar_por=marca&formato=pdf&detalle=1

Agregar un criterio nuevo es agregar una entrada al diccionario CRITERIOS.
Todos los reportes aceptan además los filtros de la búsqueda avanzada del
listado de bienes (apps/bienes/filtros.py), así que se puede pedir, por
ejemplo, "por marca, solo del Palacio Municipal, adquiridos en 2024".

Dos niveles de salida:
  - resumen  (detalle=0): una fila por grupo con cantidad y montos totalizados.
  - detallado(detalle=1): además, la lista de bienes dentro de cada grupo.
"""
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .comunes import (
    CanvasNumerado, bienes_para_reporte, cabecera_reporte, nombre_archivo,
)


# Texto que se usa cuando el bien no tiene valor en el campo por el que se
# agrupa. Se ordena siempre al final del reporte.
SIN_DATO = '(SIN DATO)'


def _texto(valor):
    """Normaliza texto libre para agrupar: sin espacios sobrantes y en mayúsculas.

    La orden de compra y la marca son campos de texto escritos a mano, así que
    'HP ' y 'hp' deben caer en el mismo grupo; sin esto el reporte saldría con
    grupos duplicados.
    """
    limpio = (valor or '').strip().upper()
    return limpio or SIN_DATO


def _denominacion(bien):
    if bien.denominacion:
        return _texto(bien.denominacion.nombre)
    return _texto(bien.descripcion)


def _cuenta(bien):
    if bien.cuenta_contable:
        return f"{bien.cuenta_contable.codigo} - {bien.cuenta_contable.descripcion}"
    return SIN_DATO


def _usuario(bien):
    if bien.usuario_asignado:
        return f"{bien.usuario_asignado.apellidos}, {bien.usuario_asignado.nombres}"
    return 'SIN ASIGNAR'


# Cada criterio: etiqueta de la columna y función que devuelve el grupo del bien.
CRITERIOS = {
    'orden_compra': {
        'label': 'Orden de Compra / Resolución de Alta',
        'titulo': 'REPORTE DE BIENES POR ORDEN DE COMPRA',
        'clave': lambda b: _texto(b.resolucion_alta),
    },
    'cuenta_contable': {
        'label': 'Cuenta Contable',
        'titulo': 'REPORTE DE BIENES POR CUENTA CONTABLE',
        'clave': _cuenta,
    },
    'denominacion': {
        'label': 'Denominación',
        'titulo': 'REPORTE DE BIENES POR DENOMINACIÓN',
        'clave': _denominacion,
    },
    'marca': {
        'label': 'Marca',
        'titulo': 'REPORTE DE BIENES POR MARCA',
        'clave': lambda b: _texto(b.marca),
    },
    'grupo_generico': {
        'label': 'Grupo Genérico',
        'titulo': 'REPORTE DE BIENES POR GRUPO GENÉRICO',
        'clave': lambda b: _texto(
            b.denominacion.grupo_generico.nombre
            if b.denominacion and b.denominacion.grupo_generico else b.grupo_generico
        ),
    },
    'clase': {
        'label': 'Clase',
        'titulo': 'REPORTE DE BIENES POR CLASE',
        'clave': lambda b: _texto(
            b.denominacion.clase.nombre if b.denominacion and b.denominacion.clase else b.clase
        ),
    },
    'local': {
        'label': 'Local',
        'titulo': 'REPORTE DE BIENES POR LOCAL',
        'clave': lambda b: _texto(b.local.nombre if b.local else None),
    },
    'area': {
        'label': 'Área',
        'titulo': 'REPORTE DE BIENES POR ÁREA',
        'clave': lambda b: _texto(b.area.nombre if b.area else None),
    },
    'oficina': {
        'label': 'Oficina',
        'titulo': 'REPORTE DE BIENES POR OFICINA',
        'clave': lambda b: _texto(b.oficina.nombre if b.oficina else None),
    },
    'usuario': {
        'label': 'Usuario Asignado',
        'titulo': 'REPORTE DE BIENES POR USUARIO ASIGNADO',
        'clave': _usuario,
    },
    'estado': {
        'label': 'Condición',
        'titulo': 'REPORTE DE BIENES POR CONDICIÓN',
        'clave': lambda b: b.get_estado_display(),
    },
    'situacion': {
        'label': 'Situación',
        'titulo': 'REPORTE DE BIENES POR SITUACIÓN (USO / DESUSO)',
        'clave': lambda b: b.get_situacion_display(),
    },
    'forma_adquisicion': {
        'label': 'Forma de Adquisición',
        'titulo': 'REPORTE DE BIENES POR FORMA DE ADQUISICIÓN',
        'clave': lambda b: b.get_forma_adquisicion_display(),
    },
    'anio_adquisicion': {
        'label': 'Año de Adquisición',
        'titulo': 'REPORTE DE BIENES POR AÑO DE ADQUISICIÓN',
        'clave': lambda b: str(b.fecha_adquisicion.year) if b.fecha_adquisicion else SIN_DATO,
    },
}

CRITERIO_POR_DEFECTO = 'cuenta_contable'


def _agrupar(bienes, criterio):
    """Agrupa la lista de bienes y calcula los totales de cada grupo.

    Devuelve una lista de dicts ordenada por nombre de grupo, con el grupo
    "(SIN DATO)" siempre al final.
    """
    funcion_clave = CRITERIOS[criterio]['clave']
    grupos = {}
    for bien in bienes:
        clave = funcion_clave(bien)
        grupo = grupos.setdefault(clave, {
            'nombre': clave,
            'bienes': [],
            'valor_adquisicion': Decimal('0.00'),
            'valor_neto': Decimal('0.00'),
        })
        grupo['bienes'].append(bien)
        grupo['valor_adquisicion'] += bien.valor_adquisicion or Decimal('0.00')
        grupo['valor_neto'] += bien.valor_neto or Decimal('0.00')

    ordenados = sorted(grupos.values(), key=lambda g: (g['nombre'] == SIN_DATO, g['nombre']))
    for grupo in ordenados:
        grupo['cantidad'] = len(grupo['bienes'])
        grupo['bienes'].sort(key=lambda b: b.codigo_patrimonial or '')
    return ordenados


def _datos_bien(bien):
    """Fila de detalle de un bien, en el mismo orden para PDF y Excel."""
    marca_modelo = ' '.join(p for p in [bien.marca, bien.modelo] if p) or '-'
    ubicacion = ' / '.join(p for p in [
        bien.local.nombre if bien.local else None,
        bien.oficina.nombre if bien.oficina else None,
    ] if p) or '-'
    return {
        'codigo': bien.codigo_patrimonial or '-',
        'denominacion': (bien.denominacion.nombre if bien.denominacion else bien.descripcion) or '-',
        'marca_modelo': marca_modelo,
        'ubicacion': ubicacion,
        'usuario': _usuario(bien),
        'fecha': bien.fecha_adquisicion.strftime('%d/%m/%Y') if bien.fecha_adquisicion else '-',
        'estado': bien.get_estado_display(),
        'valor_adquisicion': float(bien.valor_adquisicion or 0),
        'valor_neto': float(bien.valor_neto or 0),
    }


def reporte_agrupado(request):
    """Punto de entrada único de los reportes agrupados (PDF o Excel)."""
    criterio = (request.GET.get('agrupar_por') or CRITERIO_POR_DEFECTO).strip()
    if criterio not in CRITERIOS:
        criterio = CRITERIO_POR_DEFECTO
    formato = (request.GET.get('formato') or 'pdf').strip().lower()
    con_detalle = (request.GET.get('detalle') or '') in ('1', 'true', 'si', 'on')

    bienes, filtros_texto, hay_filtros = bienes_para_reporte(request)
    grupos = _agrupar(bienes, criterio)

    if formato == 'excel':
        return _excel_agrupado(criterio, grupos, filtros_texto, hay_filtros, con_detalle)
    return _pdf_agrupado(criterio, grupos, filtros_texto, hay_filtros, con_detalle)


# ==============================================================================
# SALIDA EN PDF
# ==============================================================================

def _pdf_agrupado(criterio, grupos, filtros_texto, hay_filtros, con_detalle):
    config = CRITERIOS[criterio]

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="{}"'.format(
        nombre_archivo(f"reporte_por_{criterio}", 'pdf', hay_filtros)
    )

    # El detalle lleva 10 columnas: solo entra en horizontal.
    pagesize = landscape(A4) if con_detalle else A4
    doc = SimpleDocTemplate(
        response, pagesize=pagesize,
        leftMargin=1 * cm, rightMargin=1 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm
    )
    ancho_total = pagesize[0] - 2 * cm

    styles = getSampleStyleSheet()
    th = ParagraphStyle('Th', parent=styles['Normal'], fontSize=8,
                        fontName='Helvetica-Bold', alignment=TA_CENTER)
    td = ParagraphStyle('Td', parent=styles['Normal'], fontSize=7, leading=9, alignment=TA_LEFT)
    td_num = ParagraphStyle('TdNum', parent=td, alignment=TA_RIGHT)
    td_centro = ParagraphStyle('TdCentro', parent=td, alignment=TA_CENTER)
    grupo_style = ParagraphStyle('Grupo', parent=styles['Normal'], fontSize=9,
                                 fontName='Helvetica-Bold', alignment=TA_LEFT)
    total_style = ParagraphStyle('Total', parent=styles['Normal'], fontSize=9,
                                 fontName='Helvetica-Bold', alignment=TA_RIGHT)

    titulo = config['titulo'] + (' (DETALLADO)' if con_detalle else ' (RESUMEN)')
    elements = cabecera_reporte(titulo, filtros_texto, ancho_total=ancho_total)

    total_cantidad = sum(g['cantidad'] for g in grupos)
    total_adquisicion = sum(g['valor_adquisicion'] for g in grupos)
    total_neto = sum(g['valor_neto'] for g in grupos)

    if not grupos:
        elements.append(Paragraph("No se encontraron bienes con los filtros indicados.",
                                  styles['Normal']))
        doc.build(elements, canvasmaker=CanvasNumerado)
        return response

    if con_detalle:
        elements.extend(_pdf_bloques_detalle(
            grupos, config, ancho_total, th, td, td_num, td_centro, grupo_style
        ))
    else:
        elements.append(_pdf_tabla_resumen(
            grupos, config, ancho_total, th, td, td_num, td_centro, total_cantidad
        ))

    # Total general
    elements.append(Spacer(1, 0.3 * cm))
    resumen_final = (
        f"TOTAL GENERAL &nbsp;&nbsp; Grupos: {len(grupos)} &nbsp;&nbsp; "
        f"Bienes: {total_cantidad} &nbsp;&nbsp; "
        f"Valor adquisición: S/ {total_adquisicion:,.2f} &nbsp;&nbsp; "
        f"Valor neto: S/ {total_neto:,.2f}"
    )
    tabla_total = Table([[Paragraph(resumen_final, total_style)]], colWidths=[ancho_total])
    tabla_total.setStyle(TableStyle([
        ('BOX', (0, 0), (-1, -1), 1, colors.black),
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#E8E8E8')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(tabla_total)

    doc.build(elements, canvasmaker=CanvasNumerado)
    return response


def _pdf_tabla_resumen(grupos, config, ancho_total, th, td, td_num, td_centro, total_cantidad):
    """Una fila por grupo: cantidad, montos y porcentaje sobre el total."""
    filas = [[
        Paragraph("N°", th), Paragraph(config['label'], th), Paragraph("Cant.", th),
        Paragraph("Valor Adquisición S/", th), Paragraph("Valor Neto S/", th),
        Paragraph("% Bienes", th),
    ]]
    for idx, grupo in enumerate(grupos, 1):
        porcentaje = (grupo['cantidad'] / total_cantidad * 100) if total_cantidad else 0
        filas.append([
            Paragraph(str(idx), td_centro),
            Paragraph(grupo['nombre'], td),
            Paragraph(f"{grupo['cantidad']:,}", td_centro),
            Paragraph(f"{grupo['valor_adquisicion']:,.2f}", td_num),
            Paragraph(f"{grupo['valor_neto']:,.2f}", td_num),
            Paragraph(f"{porcentaje:,.1f}%", td_num),
        ])

    proporciones = [0.06, 0.44, 0.08, 0.16, 0.16, 0.10]
    tabla = Table(filas, colWidths=[ancho_total * p for p in proporciones], repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#366092')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#999999')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F8FC')]),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    return tabla


def _pdf_bloques_detalle(grupos, config, ancho_total, th, td, td_num, td_centro, grupo_style):
    """Por cada grupo: encabezado, tabla de sus bienes y subtotal."""
    elementos = []
    encabezados = [
        Paragraph("N°", th), Paragraph("Cód. Patrimonial", th), Paragraph("Denominación", th),
        Paragraph("Marca / Modelo", th), Paragraph("Local / Oficina", th),
        Paragraph("Usuario asignado", th), Paragraph("Fec. Adq.", th), Paragraph("Cond.", th),
        Paragraph("V. Adquis. S/", th), Paragraph("V. Neto S/", th),
    ]
    proporciones = [0.035, 0.095, 0.205, 0.105, 0.135, 0.135, 0.065, 0.055, 0.085, 0.085]
    anchos = [ancho_total * p for p in proporciones]

    for grupo in grupos:
        titulo_grupo = (
            f"{config['label'].upper()}: {grupo['nombre']} &nbsp;&nbsp;—&nbsp;&nbsp; "
            f"{grupo['cantidad']} bien(es)"
        )
        cabecera_grupo = Table([[Paragraph(titulo_grupo, grupo_style)]], colWidths=[ancho_total])
        cabecera_grupo.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#CFE2FF')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#666666')),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elementos.append(cabecera_grupo)

        filas = [encabezados]
        for idx, bien in enumerate(grupo['bienes'], 1):
            d = _datos_bien(bien)
            filas.append([
                Paragraph(str(idx), td_centro),
                Paragraph(d['codigo'], td),
                Paragraph(d['denominacion'], td),
                Paragraph(d['marca_modelo'], td),
                Paragraph(d['ubicacion'], td),
                Paragraph(d['usuario'], td),
                Paragraph(d['fecha'], td_centro),
                Paragraph(d['estado'], td_centro),
                Paragraph(f"{d['valor_adquisicion']:,.2f}", td_num),
                Paragraph(f"{d['valor_neto']:,.2f}", td_num),
            ])
        filas.append([
            '', '', '', '', '', '',
            Paragraph("SUBTOTAL", th), Paragraph(f"{grupo['cantidad']:,}", th),
            Paragraph(f"{grupo['valor_adquisicion']:,.2f}", th),
            Paragraph(f"{grupo['valor_neto']:,.2f}", th),
        ])

        tabla = Table(filas, colWidths=anchos, repeatRows=1)
        tabla.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#366092')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#999999')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E8E8E8')),
            ('LEFTPADDING', (0, 0), (-1, -1), 3),
            ('RIGHTPADDING', (0, 0), (-1, -1), 3),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ]))
        elementos.append(tabla)
        elementos.append(Spacer(1, 0.35 * cm))

    return elementos


# ==============================================================================
# SALIDA EN EXCEL
# ==============================================================================

_BORDE = Border(left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin'))
_CENTRO = Alignment(horizontal='center', vertical='center')
_IZQUIERDA = Alignment(horizontal='left', vertical='center')
_DERECHA = Alignment(horizontal='right', vertical='center')


def _escribir_encabezados(ws, fila, encabezados, color="366092"):
    relleno = PatternFill(start_color=color, end_color=color, fill_type="solid")
    fuente = Font(bold=True, color="FFFFFF", size=11)
    for columna, texto in enumerate(encabezados, 1):
        celda = ws.cell(row=fila, column=columna)
        celda.value = texto
        celda.fill = relleno
        celda.font = fuente
        celda.alignment = _CENTRO
        celda.border = _BORDE


def _fila_filtros(ws, columnas, filtros_texto):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columnas)
    celda = ws.cell(row=1, column=1)
    celda.value = 'Filtros aplicados: ' + (
        ' | '.join(filtros_texto) if filtros_texto else 'ninguno (todos los bienes activos)'
    )
    celda.font = Font(bold=True, size=10)
    celda.alignment = _IZQUIERDA
    celda.fill = PatternFill(start_color="F2F6FC", end_color="F2F6FC", fill_type="solid")


def _excel_agrupado(criterio, grupos, filtros_texto, hay_filtros, con_detalle):
    config = CRITERIOS[criterio]
    wb = Workbook()

    # ---- Hoja 1: resumen por grupo ----
    ws = wb.active
    ws.title = "Resumen"
    encabezados = [config['label'], 'Cantidad', 'Valor Adquisición', 'Valor Neto', '% Bienes']
    _fila_filtros(ws, len(encabezados), filtros_texto)
    _escribir_encabezados(ws, 2, encabezados)

    total_cantidad = sum(g['cantidad'] for g in grupos)
    fila = 3
    for grupo in grupos:
        porcentaje = (grupo['cantidad'] / total_cantidad) if total_cantidad else 0
        valores = [
            grupo['nombre'], grupo['cantidad'],
            float(grupo['valor_adquisicion']), float(grupo['valor_neto']), porcentaje,
        ]
        for columna, valor in enumerate(valores, 1):
            celda = ws.cell(row=fila, column=columna)
            celda.value = valor
            celda.border = _BORDE
            if columna == 1:
                celda.alignment = _IZQUIERDA
            elif columna == 2:
                celda.alignment = _CENTRO
            elif columna == 5:
                celda.number_format = '0.0%'
                celda.alignment = _DERECHA
            else:
                celda.number_format = '#,##0.00'
                celda.alignment = _DERECHA
        fila += 1

    # Fila de total general
    total_adquisicion = sum(float(g['valor_adquisicion']) for g in grupos)
    total_neto = sum(float(g['valor_neto']) for g in grupos)
    for columna, valor in enumerate(['TOTAL GENERAL', total_cantidad,
                                     total_adquisicion, total_neto, 1 if grupos else 0], 1):
        celda = ws.cell(row=fila, column=columna)
        celda.value = valor
        celda.font = Font(bold=True)
        celda.border = _BORDE
        celda.fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
        if columna in (3, 4):
            celda.number_format = '#,##0.00'
            celda.alignment = _DERECHA
        elif columna == 5:
            celda.number_format = '0.0%'
            celda.alignment = _DERECHA
        elif columna == 2:
            celda.alignment = _CENTRO

    for columna, ancho in zip('ABCDE', [55, 12, 20, 18, 12]):
        ws.column_dimensions[columna].width = ancho
    ws.freeze_panes = 'A3'

    # ---- Hoja 2: detalle de bienes (una sola hoja con la columna del grupo) ----
    # Se usa una hoja única en vez de una por grupo: los nombres de grupo son
    # texto libre (marcas, órdenes de compra) y Excel no admite varios de los
    # caracteres que traen ni nombres de hoja repetidos o de más de 31 letras.
    if con_detalle:
        ws_det = wb.create_sheet("Detalle")
        encabezados_det = [
            config['label'], 'Cód. Patrimonial', 'Denominación', 'Marca / Modelo',
            'Local / Oficina', 'Usuario asignado', 'Fecha Adquisición', 'Condición',
            'Valor Adquisición', 'Valor Neto',
        ]
        _fila_filtros(ws_det, len(encabezados_det), filtros_texto)
        _escribir_encabezados(ws_det, 2, encabezados_det)

        fila = 3
        for grupo in grupos:
            for bien in grupo['bienes']:
                d = _datos_bien(bien)
                valores = [
                    grupo['nombre'], d['codigo'], d['denominacion'], d['marca_modelo'],
                    d['ubicacion'], d['usuario'], d['fecha'], d['estado'],
                    d['valor_adquisicion'], d['valor_neto'],
                ]
                for columna, valor in enumerate(valores, 1):
                    celda = ws_det.cell(row=fila, column=columna)
                    celda.value = valor
                    celda.border = _BORDE
                    if columna in (9, 10):
                        celda.number_format = '#,##0.00'
                        celda.alignment = _DERECHA
                    elif columna in (7, 8):
                        celda.alignment = _CENTRO
                    else:
                        celda.alignment = _IZQUIERDA
                fila += 1

        for columna, ancho in zip(['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J'],
                                  [40, 18, 40, 25, 35, 30, 16, 12, 18, 15]):
            ws_det.column_dimensions[columna].width = ancho
        ws_det.freeze_panes = 'A3'
        # El autofiltro permite al usuario recortar el detalle sin volver a pedir
        # el reporte al servidor.
        ws_det.auto_filter.ref = f"A2:J{max(fila - 1, 2)}"

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="{}"'.format(
        nombre_archivo(f"reporte_por_{criterio}", 'xlsx', hay_filtros)
    )
    wb.save(response)
    return response
